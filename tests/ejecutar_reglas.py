"""Ejecuta las reglas de consistencia de tests/reglas_consistencia.xlsx contra la tabla
personas (vista 'personas' de app/db.py). Solo lectura, en el propio proceso (no hace
falta el servidor).

    python tests/ejecutar_reglas.py A        # bloque A, B o C
    python tests/ejecutar_reglas.py A 20 60  # solo las reglas 20..60 del bloque

Cada regla es una condicion SQL que selecciona las filas que la INCUMPLEN: el test pasa
si devuelve 0 filas. Escribe tests/resultados/<bloque>.csv (version "actual", la que leen
construir_resultados.py y demas scripts) con el nº de filas incumplidoras, el % sobre el
total, hasta 5 IPF de ejemplo y el tiempo - y ADEMAS una copia en
tests/resultados/historico/<fecha_hora>/<bloque>.csv (2026-09-22, decision explicita del
usuario: conservar los resultados sucesivos segun se van arreglando errores, para poder
comparar una ejecucion con la anterior en vez de perder el resultado viejo al sobrescribir).
"""
import csv, os, shutil, sys, time
from datetime import datetime

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RAIZ)
from openpyxl import load_workbook
from app.db import get_connection, PARQUET_DIR

bloque = sys.argv[1].upper()
desde = int(sys.argv[2]) if len(sys.argv) > 2 else 1
hasta = int(sys.argv[3]) if len(sys.argv) > 3 else 10 ** 6

import glob
# la version con mayor nombre (reglas_consistencia_v2 > reglas_consistencia), no la mas reciente por fecha:
# un fichero abierto y guardado desde Excel cambiaria su fecha y se leeria una version antigua sin querer
FICHERO = max(glob.glob(os.path.join(RAIZ, "tests", "reglas_consistencia*.xlsx")), key=os.path.basename)
print("reglas leidas de:", os.path.basename(FICHERO))
ws = load_workbook(FICHERO, read_only=True)["Reglas"]
reglas = [r for r in ws.iter_rows(min_row=2, values_only=True)
          if r[0] and r[0].startswith(bloque + "-") and desde <= int(r[0].split("-")[1]) <= hasta]

con = get_connection()
# Tabla auxiliar para las reglas del grupo G08 (2026-09-24, decision explicita del usuario): personas
# cuyos episodios de afiliacion son TODOS de prestacion/subsidio (TRL 700-799) o TRL 400, sin ningun
# episodio de empleo real. Es dato de origen (p. ej. agrarios con trabajo temporal mal registrado),
# no un fallo del ETL: sus campos de empleo quedan a NULL a proposito, asi que esas reglas las excluyen.
if any("solo_prestaciones" in (r[8] or "") for r in reglas):
    con.execute("""CREATE TEMP TABLE solo_prestaciones AS
        SELECT ipf FROM afiliacion GROUP BY ipf
        HAVING bool_and(COALESCE(trl >= 700 AND trl < 800, FALSE) OR COALESCE(trl = 400, FALSE))""")
# Menores autorizados (2026-09-24, G04, decision explicita del usuario): personas con algun episodio de
# afiliacion que EMPIEZA antes de la edad minima legal (14 hasta la cohorte 1965, 16 despues) en un regimen o
# sector donde es plausible que trabajen menores con autorizacion: el Estatuto de los Trabajadores (art. 6.4)
# permite la intervencion de menores de 16 en espectaculos publicos. Regimen 0112 = Regimen General (Artistas)
# (guia MCVL 2.001); CNAE 2009 (3 digitos): 59 cine/video/TV, 60 radio y TV, 90 creacion artistica y espectaculos,
# 93 deportivas/recreativas, 73 publicidad, 781 agencias de colocacion (incluye casting), 742 fotografia,
# 855 educacion cultural (musica, danza, arte dramatico). CNAE desconocido (0) NO cuenta como plausible.
if any("menor_autorizado" in (r[8] or "") for r in reglas):
    con.execute("""CREATE TEMP TABLE menor_autorizado AS
        SELECT DISTINCT a.ipf
        FROM afiliacion a JOIN personas p ON p.ipf = a.ipf
        WHERE date_diff('year', p.fecha_nacimiento, a.fecha_alta_real_afil)
                < (CASE WHEN year(p.fecha_nacimiento) <= 1965 THEN 14 ELSE 16 END)
          AND (a.reg_cotizacion = 112
               OR a.cnae2009_cc // 10 IN (59, 60, 90, 93, 73)
               OR a.cnae2009_cc IN (781, 742, 855))""")
# Personas con alguna pension de orfandad VITALICIA (2026-09-24, C-153): segun la guia MCVL (4.030) es la
# orfandad de quien tiene incapacidad permanente reconocida, que no tiene limite de edad (art. 224 LGSS).
# La vista "pensiones" de app/db.py no expone esa columna, asi que se lee directamente del Parquet.
if any("orfandad_vitalicia" in (r[8] or "") for r in reglas):
    _pens = (PARQUET_DIR / "tabla_pensiones_extendida.parquet").as_posix()
    con.execute(f"""CREATE TEMP TABLE orfandad_vitalicia AS
        SELECT DISTINCT "01 IPF" AS ipf FROM read_parquet('{_pens}')
        WHERE "Tipo_pension" = 'Orfandad' AND "30 PREST_VITALICIA (ORFANDAD Y VIUDEDAD)" = 2""")
total = con.execute("SELECT COUNT(*) FROM personas").fetchone()[0]
filas, t_bloque = [], time.time()
for r in reglas:
    id_, familia, regla, conf, sql = r[0], r[2], r[5], r[7], r[8]
    t0 = time.time()
    try:
        n = con.execute(f"SELECT COUNT(*) FROM personas WHERE {sql}").fetchone()[0]
        muestra = ""
        if n:
            muestra = ",".join(str(x[0]) for x in con.execute(f"SELECT ipf FROM personas WHERE {sql} ORDER BY ipf LIMIT 5").fetchall())
        err = ""
    except Exception as e:
        n, muestra, err = None, "", str(e).splitlines()[0][:200]
    filas.append([id_, familia, regla, conf.split()[0], "" if n is None else n,
                  "" if n is None else round(100 * n / total, 4), muestra, round(time.time() - t0, 2), err])

os.makedirs(os.path.join(RAIZ, "tests", "resultados"), exist_ok=True)
nombre = f"{bloque}.csv" if desde == 1 and hasta == 10 ** 6 else f"{bloque}_{desde}_{hasta}.csv"
ruta_actual = os.path.join(RAIZ, "tests", "resultados", nombre)
with open(ruta_actual, "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f, delimiter=";")
    w.writerow(["ID", "Familia", "Regla", "Confianza", "N_incumplen", "Pct_incumplen", "Ejemplo_ipf", "Segundos", "Error"])
    w.writerows(filas)

# Copia con fecha y hora, ver docstring - conserva el historico de ejecuciones sucesivas
# (una por cada vez que se corrigen errores y se regeneran las tablas) sin tocar la version
# "actual" de arriba, que es la que siguen leyendo construir_resultados.py y agrupar.py.
carpeta_historico = os.path.join(RAIZ, "tests", "resultados", "historico", datetime.now().strftime("%Y-%m-%d_%H%M"))
os.makedirs(carpeta_historico, exist_ok=True)
shutil.copy(ruta_actual, os.path.join(carpeta_historico, nombre))

ok = [x for x in filas if x[4] == 0]
ko = sorted([x for x in filas if x[4] not in (0, "")], key=lambda x: -x[4])
er = [x for x in filas if x[4] == ""]
print(f"===== bloque {bloque}: {len(filas)} reglas en {round(time.time() - t_bloque)} s sobre {total:,} personas =====")
print(f"pasan (0 filas): {len(ok)} | incumplidas: {len(ko)} | con error de ejecución: {len(er)}")
for x in ko:
    print(f"  {x[0]} [{x[3]}] {x[4]:>9,} ({x[5]:.3f}%)  {x[2][:110]}")
for x in er:
    print(f"  ERROR {x[0]}: {x[8]}")
