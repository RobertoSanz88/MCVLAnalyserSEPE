"""Prueba de humo por bloques: python tests/smoke.py <bloque>
bloques: colectivos | cruce_a | cruce_b | cruce_extra | trayectorias

Requiere el servidor ya arrancado en http://127.0.0.1:8000 (uvicorn app.main:app
--port 8000) y los Parquet reales cargados. Solo hace consultas de lectura.
Cada bloque cabe en el limite de 10 min por comando. Imprime los fallos (HTTP
inesperado o invariante rota) y las llamadas de mas de 15 s.

Primera ejecucion (2026-09-21), 5 bloques / 277 llamadas / ~7 min: 0 fallos en
colectivos, cruce_a, cruce_b y cruce_extra; 2 fallos en trayectorias (500 en
/trayectorias/<ipf inexistente>), corregidos en _persona_resumen (ahora 404).
No cubre las trayectorias de afiliacion (40-110 s por consulta).
"""
import io, random, sys, time
import requests
from openpyxl import load_workbook

B = "http://127.0.0.1:8000"
S = requests.Session()
fails, t_slow = [], []
n_calls = 0


def call(name, method, path, expect=200, timeout=300, **kw):
    global n_calls
    n_calls += 1
    t0 = time.time()
    try:
        r = S.request(method, B + path, timeout=timeout, **kw)
    except Exception as e:
        fails.append((name, f"EXCEPCION {type(e).__name__}: {e}"))
        return None
    dt = time.time() - t0
    if dt > 15:
        t_slow.append((name, round(dt, 1)))
    ok = r.status_code == expect if isinstance(expect, int) else r.status_code in expect
    if not ok:
        fails.append((name, f"HTTP {r.status_code} (esperado {expect}): {r.text[:250]}"))
        return None
    return r


def bad(name, msg):
    fails.append((name, msg))


def cat(col, vals):
    return {"columna": col, "tipo": "categorica", "categorias": vals}


def num(col, op, v):
    return {"columna": col, "tipo": "numerica", "operador": op, "valor": v}


def xlsx_ok(name, rr):
    if rr is None:
        return
    try:
        wb = load_workbook(io.BytesIO(rr.content))
        if not wb.sheetnames:
            bad(name, "xlsx sin hojas")
    except Exception as e:
        bad(name, f"xlsx invalido: {e}")


cols = call("columnas", "GET", "/colectivos/columnas").json()
NUMS, CATS = cols["numericas"], cols["categoricas"]
EDAD = "edad_actual_fecha_corte"
FILTROS = {
    "cat_cat": dict(filtro1=cat("sexo", ["Hombre"]), filtro2=cat("nacionalidad_agrupado", ["espana"])),
    "num_cat": dict(filtro1=num(EDAD, ">=", 45), filtro2=cat("sexo", ["Mujer"])),
    "num_num": dict(filtro1=num(EDAD, "<", 30), filtro2=num("duracion_promedio_empleo", ">", 500)),
    "con_f3f4": dict(filtro1=cat("sexo", ["Hombre"]), filtro2=cat("nacionalidad_agrupado", ["espana"]),
                     filtro3=cat("situacion_paro_ocupacion", ["Ocupado"]), filtro4=num(EDAD, ">=", 50)),
    "solapado": dict(filtro1=cat("sexo", ["Hombre"]), filtro2=cat("sexo", ["Hombre"])),
    "vacio_sexo_opuesto": dict(filtro1=cat("sexo", ["Hombre"]), filtro2=cat("sexo", ["Mujer"])),
    "multi_valor": dict(filtro1=cat("situacion_paro_ocupacion", ["Ocupado", "Parado_LD"]), filtro2=num(EDAD, ">=", 18)),
    "umbral_extremo": dict(filtro1=num(EDAD, ">=", 999), filtro2=cat("sexo", ["Mujer"])),
    "no_consta": dict(filtro1=cat("act_econ_ppal", ["NO CONSTA"]), filtro2=cat("sexo", ["Hombre"])),
}
F0 = FILTROS["cat_cat"]
ORD = ["filtro1_si__filtro2_si", "filtro1_si__filtro2_no", "filtro1_no__filtro2_si", "filtro1_no__filtro2_no"]


# ---------------- invariantes de cruce ----------------
def inv_catcat(name, d):
    for k, sub in d["subgrupos"].items():
        x, y, z = sub["x"], sub["y"], sub["z"]
        if len(z) != len(y) or any(len(f) != len(x) for f in z):
            return bad(name, f"{k}: forma de z incoherente con x/y")
        if any(v < 0 for f in z for v in f):
            return bad(name, f"{k}: z con negativos")
        if sum(sum(f) for f in z) > sub["recuento"]:
            return bad(name, f"{k}: suma(z) > recuento ({sum(sum(f) for f in z)} > {sub['recuento']})")
        if len(set(x)) != len(x) or len(set(y)) != len(y):
            return bad(name, f"{k}: categorias duplicadas en un eje")


def inv_numcat(name, d):
    for k, sub in d["subgrupos"].items():
        for c in sub["cajas"]:
            if not c["n"] > 0:
                return bad(name, f"{k}: caja con n<=0 en {c['categoria']}")
            if c["q1"] is not None:
                if not (c["min"] <= c["q1"] <= c["mediana"] <= c["q3"] <= c["max"]):
                    return bad(name, f"{k}: cuartiles desordenados en {c['categoria']}: {c}")
                if not (c["min"] <= c["bigote_inf"] <= c["q1"] and c["q3"] <= c["bigote_sup"] <= c["max"]):
                    return bad(name, f"{k}: bigotes fuera de rango en {c['categoria']}: {c}")
        if sub["histograma_bins"]:
            ns = {c["categoria"]: c["n"] for c in sub["cajas"]}
            for cat_, h in sub["histograma_por_categoria"].items():
                if sum(h) != ns.get(cat_, 0):
                    return bad(name, f"{k}: suma histograma {sum(h)} != n {ns.get(cat_, 0)} en {cat_}")
        if sum(c["n"] for c in sub["cajas"]) > sub["recuento"]:
            return bad(name, f"{k}: suma(n cajas) > recuento")


def inv_numnum(name, d):
    for k, sub in d["subgrupos"].items():
        z = sub["z"]
        if sub["x"] and (len(z) != 20 or any(len(f) != 20 for f in z)):
            return bad(name, f"{k}: rejilla no 20x20")
        if sum(sum(f) for f in z) > sub["recuento"]:
            return bad(name, f"{k}: suma(z) > recuento")


INV = {"categorica_categorica": inv_catcat, "numerica_categorica": inv_numcat, "numerica_numerica": inv_numnum}


def cruce(name, filtros, a, b):
    rr = call(name, "POST", "/colectivos/cruce", json={**filtros, "variable1": a, "variable2": b})
    if not rr:
        return None
    d = rr.json()
    if d["variable1"] != a or d["variable2"] != b:
        bad(name, f"variables reordenadas: {d['variable1']},{d['variable2']} != {a},{b}")
    INV[d["tipo"]](name, d)
    return d


def pares_cruce():
    random.seed(7)
    pares = []
    for c in CATS:
        if c != "sexo": pares.append((c, "sexo"))
        if c != "act_econ_ppal": pares.append((c, "act_econ_ppal"))
    for n_ in NUMS:
        pares += [(n_, "sexo"), ("sexo", n_), (n_, "act_econ_ppal")]
        if n_ != EDAD: pares.append((n_, EDAD))
    random.shuffle(pares)
    return pares[:110]


# ---------------- bloques ----------------
def b_colectivos():
    for p in ["/", "/colectivos", "/health"]:
        call(f"GET {p}", "GET", p)
    for c in CATS:
        rr = call(f"valores/{c}", "GET", f"/colectivos/valores/{c}")
        if rr:
            v = rr.json()["valores"]
            if not v: bad(f"valores/{c}", "lista vacia")
            if any(x is None for x in v): bad(f"valores/{c}", "contiene None")
    call("valores inyeccion", "GET", "/colectivos/valores/sexo;DROP", expect=400)
    call("valores inexistente", "GET", "/colectivos/valores/no_existe", expect=400)
    call("valores numerica", "GET", f"/colectivos/valores/{NUMS[0]}", expect=400)
    call("filtro cat sin categorias", "POST", "/colectivos/subgrupos", expect=400,
         json=dict(filtro1=cat("sexo", []), filtro2=cat("sexo", ["Mujer"])))
    call("filtro columna inexistente", "POST", "/colectivos/subgrupos", expect=400,
         json=dict(filtro1=cat("no_existe", ["a"]), filtro2=cat("sexo", ["Mujer"])))
    call("filtro columna inyeccion", "POST", "/colectivos/subgrupos", expect=400,
         json=dict(filtro1=cat("sexo) OR 1=1 --", ["a"]), filtro2=cat("sexo", ["Mujer"])))
    call("filtro operador invalido", "POST", "/colectivos/subgrupos", expect=422,
         json=dict(filtro1={"columna": EDAD, "tipo": "numerica", "operador": "LIKE", "valor": 1}, filtro2=cat("sexo", ["Mujer"])))
    call("filtro sin filtro2", "POST", "/colectivos/subgrupos", expect=422, json=dict(filtro1=cat("sexo", ["Mujer"])))
    for fname, f in FILTROS.items():
        rr = call(f"recuentos[{fname}]", "POST", "/colectivos/recuentos", json=f)
        if rr:
            d = rr.json()
            if any(v["recuento"] < 0 for v in d.values()): bad(f"recuentos[{fname}]", "recuento negativo")
            if fname in ("cat_cat", "num_cat", "num_num") and sum(v["recuento"] for v in d.values()) == 0:
                bad(f"recuentos[{fname}]", "los 4 subgrupos vacios")
        rr = call(f"subgrupos[{fname}]", "POST", "/colectivos/subgrupos", json=f)
        if rr:
            sg = rr.json()["subgrupos"]
            for k in ORD:
                if k not in sg: bad(f"subgrupos[{fname}]", f"falta subgrupo {k}")
        xlsx_ok(f"exportar[{fname}]", call(f"exportar[{fname}]", "POST", "/colectivos/exportar", json=f))
    call("subgrupos-sin-filtro", "GET", "/colectivos/subgrupos-sin-filtro")
    xlsx_ok("exportar-sin-filtro", call("exportar-sin-filtro", "GET", "/colectivos/exportar-sin-filtro"))


def b_cruce_a():
    for a, b in pares_cruce()[:55]:
        cruce(f"cruce {a} x {b}", F0, a, b)


def b_cruce_b():
    for a, b in pares_cruce()[55:]:
        cruce(f"cruce {a} x {b}", F0, a, b)


def b_cruce_extra():
    for fname, f in FILTROS.items():
        for a, b in [("grupo_cotiz_ppal", "act_econ_ppal"), (EDAD, "sexo"), (EDAD, "duracion_promedio_empleo")]:
            cruce(f"cruce[{fname}] {a} x {b}", f, a, b)
    call("cruce misma variable", "POST", "/colectivos/cruce", expect=400, json={**F0, "variable1": "sexo", "variable2": "sexo"})
    call("cruce var inexistente", "POST", "/colectivos/cruce", expect=400, json={**F0, "variable1": "sexo", "variable2": "zzz"})
    call("cruce var inyeccion", "POST", "/colectivos/cruce", expect=400, json={**F0, "variable1": "sexo", "variable2": "sexo;DROP"})
    call("cruce ipf", "POST", "/colectivos/cruce", expect=400, json={**F0, "variable1": "ipf", "variable2": EDAD})
    for a, b in [("prov_trabajo_ppal", "act_econ_ppal"), ("sexo", "grupo_cotiz_ppal"), ("act_econ_ppal", "prov_fiscal_actual")]:
        for v in ["ingresos_bruto_total_2024", EDAD, "duracion_promedio_empleo"]:
            name = f"central {a} x {b} -> {v}"
            rr = call(name, "POST", "/colectivos/cruce-central", json=dict(variable1=a, variable2=b, variable_central=v))
            if rr:
                d = rr.json()
                X, Y = d["x"], d["y"]
                for key in ("z_media", "z_mediana", "n"):
                    m = d[key]
                    if len(m) != len(Y) or any(len(f) != len(X) for f in m):
                        bad(name, f"{key} forma incoherente"); break
                else:
                    for i in range(len(Y)):
                        for j in range(len(X)):
                            if (d["n"][i][j] == 0) != (d["z_media"][i][j] is None):
                                bad(name, f"celda ({i},{j}) n={d['n'][i][j]} media={d['z_media'][i][j]}")
                                break
    call("central var no numerica", "POST", "/colectivos/cruce-central", expect=400, json=dict(variable1="sexo", variable2="act_econ_ppal", variable_central="sexo"))
    call("central var1 numerica", "POST", "/colectivos/cruce-central", expect=400, json=dict(variable1=EDAD, variable2="act_econ_ppal", variable_central="ingresos_bruto_total_2024"))
    call("central repetidas", "POST", "/colectivos/cruce-central", expect=400, json=dict(variable1="sexo", variable2="sexo", variable_central=EDAD))
    call("puntos", "POST", "/colectivos/cruce/puntos", json={**F0, "variable1": EDAD, "variable2": "duracion_promedio_empleo", "subgrupo": "filtro1_si__filtro2_no"})
    call("puntos subgrupo malo", "POST", "/colectivos/cruce/puntos", expect=400, json={**F0, "variable1": EDAD, "variable2": "duracion_promedio_empleo", "subgrupo": "xx"})
    call("puntos no numericas", "POST", "/colectivos/cruce/puntos", expect=400, json={**F0, "variable1": "sexo", "variable2": EDAD, "subgrupo": "filtro1_si__filtro2_no"})

    def exp(name, body, expect=200):
        rr = call(name, "POST", "/colectivos/cruce/exportar", expect=expect, json=body)
        if expect == 200:
            xlsx_ok(name, rr)

    for tg in ["contingencia", "apiladas", "apiladas100", "burbujas"]:
        exp(f"exp catcat {tg}", {"tipo": "categorica_categorica", "variable1": "grupo_cotiz_ppal", "variable2": "act_econ_ppal", "tipo_grafico": tg, **F0})
    for tg in ["boxplot", "histograma"]:
        exp(f"exp numcat {tg}", {"tipo": "numerica_categorica", "variable1": "ingresos_bruto_total_2024", "variable2": "grupo_cotiz_ppal", "tipo_grafico": tg, **F0})
        exp(f"exp catnum {tg}", {"tipo": "numerica_categorica", "variable1": "grupo_cotiz_ppal", "variable2": "ingresos_bruto_total_2024", "tipo_grafico": tg, **F0})
    d = cruce("cruce para exportar numnum", F0, EDAD, "duracion_promedio_empleo")
    if d:
        exp("exp numnum", {"tipo": "numerica_numerica", "variable1": EDAD, "variable2": "duracion_promedio_empleo", "tipo_grafico": "densidad", "datos": d["subgrupos"]})
    exp("exp central", {"tipo": "categorica_categorica_central", "variable1": "prov_trabajo_ppal", "variable2": "act_econ_ppal", "variable_central": EDAD, "tipo_grafico": "heatmap"})
    exp("exp catcat sin filtros", {"tipo": "categorica_categorica", "variable1": "sexo", "variable2": "act_econ_ppal", "tipo_grafico": "contingencia"}, expect=400)
    exp("exp numnum sin datos", {"tipo": "numerica_numerica", "variable1": EDAD, "variable2": "duracion_promedio_empleo", "tipo_grafico": "densidad"}, expect=400)
    exp("exp tipo invalido", {"tipo": "xx", "variable1": "a", "variable2": "b", "tipo_grafico": "c"}, expect=422)


def b_trayectorias():
    for fname in ["cat_cat", "num_cat", "num_num", "con_f3f4", "vacio_sexo_opuesto", "umbral_extremo"]:
        for ep in ["bases", "pensiones"]:
            rr = call(f"tray {ep}[{fname}]", "POST", f"/trayectorias/colectivo/{ep}", json=FILTROS[fname])
            if rr:
                for k, sub in rr.json().items():
                    for eje, serie in sub["series"].items():
                        xs = [p["x"] for p in serie]
                        if xs != sorted(xs): bad(f"tray {ep}[{fname}]", f"{k}/{eje}: x no ordenado")
                        for p in serie:
                            if p["n"] <= 0: bad(f"tray {ep}[{fname}]", f"{k}/{eje}: n<=0 en x={p['x']}")
                            if p["n"] == 1 and p["media"] is not None and abs(p["media"] - p["mediana"]) > 1e-6:
                                bad(f"tray {ep}[{fname}]", f"{k}/{eje}: n=1 con media!=mediana en x={p['x']}")
    for ep in ["bases-sin-filtro", "pensiones-sin-filtro"]:
        call(f"tray {ep}", "GET", f"/trayectorias/colectivo/{ep}")
    call("tray categorica tipo_pension", "POST", "/trayectorias/colectivo/categorica", json={**F0, "variable": "tipo_pension"})
    call("tray categorica variable invalida", "POST", "/trayectorias/colectivo/categorica", expect=422, json={**F0, "variable": "zzz"})
    for _ in range(6):
        rr = call("individual aleatorio", "POST", "/trayectorias/individual/aleatorio", json={})
        if rr:
            d = rr.json()
            for e in d["episodios"]:
                if e.get("inicio") and e.get("fin") and e["fin"] < e["inicio"]:
                    bad("individual aleatorio", f"IPF {d['ipf']}: episodio con fin < inicio {e}")
            call(f"pagina trayectoria {d['ipf']}", "GET", f"/trayectorias/{d['ipf']}")
    for sg in ORD:
        call(f"individual aleatorio {sg}", "POST", "/trayectorias/individual/aleatorio", json={"colectivo": F0, "subgrupo": sg}, expect=(200, 404))
    call("individual ipf inexistente", "GET", "/trayectorias/999999999", expect=(404, 200))
    call("individual ipf texto", "GET", "/trayectorias/abc", expect=(404, 422))
    call("individual ipf negativo", "GET", "/trayectorias/-5", expect=(404, 422, 200))
    call("individual colectivo null", "POST", "/trayectorias/individual/aleatorio", expect=(200, 422), json={"colectivo": None})


BLOQUES = {"colectivos": b_colectivos, "cruce_a": b_cruce_a, "cruce_b": b_cruce_b,
           "cruce_extra": b_cruce_extra, "trayectorias": b_trayectorias}

if __name__ == "__main__":
    nombre = sys.argv[1]
    t0 = time.time()
    BLOQUES[nombre]()
    print(f"===== bloque {nombre}: {n_calls} llamadas, {round(time.time()-t0)}s, {len(fails)} fallos =====")
    for n_, m in fails:
        print(" -", n_, "->", m)
    if t_slow:
        print("lentas (>15s):", t_slow)
